from __future__ import annotations

import csv
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation


def build_supported_fixtures(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Newton second law F equals m a")
    pdf.save(root / "lesson.pdf")
    pdf.close()

    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("牛顿第二定律说明力与加速度的关系。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "概念"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "力"
    table.cell(1, 1).text = "质量乘以加速度"
    document.save(root / "lesson.docx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "牛顿第二定律"
    slide.placeholders[1].text = "力等于质量乘以加速度"
    presentation.save(root / "lesson.pptx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "物理"
    sheet.append(["概念", "说明"])
    sheet.append(["力", "质量乘以加速度"])
    workbook.save(root / "lesson.xlsx")

    with (root / "lesson.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["概念", "说明"])
        writer.writerow(["力", "质量乘以加速度"])

    (root / "lesson.txt").write_text(
        "牛顿第二定律\n\n力等于质量乘以加速度。",
        encoding="utf-8",
    )
    (root / "lesson.md").write_text(
        "# 第一章\n\n牛顿第二定律\n\n## 公式\n\nF = ma",
        encoding="utf-8",
    )
    return root


def build_encrypted_pdf(path: Path) -> Path:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "protected")
    pdf.save(
        path,
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner-secret",
        user_pw="user-secret",
    )
    pdf.close()
    return path
