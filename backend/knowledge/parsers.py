from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
import re
from typing import Literal
import zipfile

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


LocatorType = Literal["page", "slide", "sheet_rows", "paragraph"]
_OFFICE_EXTENSIONS = {".docx", ".pptx", ".xlsx"}
_SUPPORTED_EXTENSIONS = _OFFICE_EXTENSIONS | {".pdf", ".csv", ".txt", ".md", ".markdown"}


class KnowledgeParseError(RuntimeError):
    def __init__(self, code: str, *, retryable: bool = False):
        super().__init__(code)
        self.code = code
        self.retryable = retryable


@dataclass(frozen=True)
class ParserLimits:
    max_text_characters: int = 50_000_000
    max_blocks: int = 100_000
    max_pdf_pages: int = 2_000
    max_zip_entries: int = 10_000
    max_uncompressed_bytes: int = 500_000_000
    max_compression_ratio: float = 100.0
    max_sheets: int = 100
    max_cells: int = 1_000_000
    max_csv_rows: int = 500_000
    max_csv_columns: int = 1_024


@dataclass(frozen=True)
class StructuredBlock:
    text: str
    heading_path: tuple[str, ...]
    locator_type: LocatorType
    locator_start: int
    locator_end: int
    sheet_name: str | None = None


@dataclass(frozen=True)
class ParsedDocument:
    blocks: tuple[StructuredBlock, ...]
    text_characters: int
    page_count: int | None = None
    slide_count: int | None = None
    sheet_count: int | None = None
    ocr_required: bool = False


class _BlockCollector:
    def __init__(self, limits: ParserLimits):
        self.limits = limits
        self.blocks: list[StructuredBlock] = []
        self.characters = 0

    def add(self, block: StructuredBlock) -> None:
        text = block.text.strip()
        if not text:
            return
        if len(self.blocks) >= self.limits.max_blocks:
            raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
        next_total = self.characters + len(text)
        if next_total > self.limits.max_text_characters:
            raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
        self.characters = next_total
        self.blocks.append(
            StructuredBlock(
                text=text,
                heading_path=block.heading_path,
                locator_type=block.locator_type,
                locator_start=block.locator_start,
                locator_end=block.locator_end,
                sheet_name=block.sheet_name,
            )
        )

    def result(
        self,
        *,
        page_count: int | None = None,
        slide_count: int | None = None,
        sheet_count: int | None = None,
        ocr_required: bool = False,
    ) -> ParsedDocument:
        return ParsedDocument(
            blocks=tuple(self.blocks),
            text_characters=self.characters,
            page_count=page_count,
            slide_count=slide_count,
            sheet_count=sheet_count,
            ocr_required=ocr_required,
        )


def _validate_signature(path: Path, extension: str) -> None:
    with path.open("rb") as stream:
        header = stream.read(5)
    if extension == ".pdf" and header != b"%PDF-":
        raise KnowledgeParseError("KNOWLEDGE_FILE_SIGNATURE_MISMATCH")
    if extension in _OFFICE_EXTENSIONS and not zipfile.is_zipfile(path):
        raise KnowledgeParseError("KNOWLEDGE_FILE_SIGNATURE_MISMATCH")


def _validate_office_archive(path: Path, limits: ParserLimits) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_zip_entries:
                raise KnowledgeParseError("KNOWLEDGE_ARCHIVE_UNSAFE")
            total_uncompressed = 0
            for info in infos:
                member = PurePosixPath(info.filename)
                if (
                    member.is_absolute()
                    or ".." in member.parts
                    or "\\" in info.filename
                    or info.flag_bits & 0x1
                    or info.filename.lower().endswith("vbaproject.bin")
                ):
                    raise KnowledgeParseError("KNOWLEDGE_ARCHIVE_UNSAFE")
                total_uncompressed += info.file_size
                if total_uncompressed > limits.max_uncompressed_bytes:
                    raise KnowledgeParseError("KNOWLEDGE_ARCHIVE_UNSAFE")
                if info.file_size:
                    if info.compress_size <= 0:
                        raise KnowledgeParseError("KNOWLEDGE_ARCHIVE_UNSAFE")
                    if info.file_size / info.compress_size > limits.max_compression_ratio:
                        raise KnowledgeParseError("KNOWLEDGE_ARCHIVE_UNSAFE")
                if info.filename.lower().endswith((".xml", ".rels")):
                    with archive.open(info) as stream:
                        prefix = stream.read(65_536).upper()
                    if b"<!DOCTYPE" in prefix or b"<!ENTITY" in prefix:
                        raise KnowledgeParseError("KNOWLEDGE_ARCHIVE_UNSAFE")
    except KnowledgeParseError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise KnowledgeParseError("KNOWLEDGE_PARSE_FAILED") from exc


def _parse_pdf(path: Path, limits: ParserLimits) -> ParsedDocument:
    collector = _BlockCollector(limits)
    try:
        document = fitz.open(path)
    except Exception as exc:
        raise KnowledgeParseError("KNOWLEDGE_PARSE_FAILED") from exc
    try:
        if document.needs_pass:
            raise KnowledgeParseError("KNOWLEDGE_DOCUMENT_ENCRYPTED")
        page_count = len(document)
        if page_count > limits.max_pdf_pages:
            raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
        for index, page in enumerate(document, 1):
            collector.add(
                StructuredBlock(
                    text=page.get_text("text"),
                    heading_path=(),
                    locator_type="page",
                    locator_start=index,
                    locator_end=index,
                )
            )
        average = collector.characters / max(page_count, 1)
        return collector.result(
            page_count=page_count,
            ocr_required=page_count > 0 and average < 20,
        )
    finally:
        document.close()


def _heading_level(style_name: str | None) -> int | None:
    matched = re.match(r"^Heading\s+([1-9])$", style_name or "", flags=re.IGNORECASE)
    return int(matched.group(1)) if matched else None


def _parse_docx(path: Path, limits: ParserLimits) -> ParsedDocument:
    collector = _BlockCollector(limits)
    document = Document(path)
    headings: list[str] = []
    paragraph_index = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        level = _heading_level(getattr(paragraph.style, "name", None))
        if level is not None:
            headings = headings[: level - 1]
            headings.append(text)
            continue
        paragraph_index += 1
        collector.add(
            StructuredBlock(
                text=text,
                heading_path=tuple(headings),
                locator_type="paragraph",
                locator_start=paragraph_index,
                locator_end=paragraph_index,
            )
        )
    for table in document.tables:
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ]
        text = "\n".join(row for row in rows if row.strip(" |"))
        if text:
            paragraph_index += 1
            collector.add(
                StructuredBlock(
                    text=text,
                    heading_path=tuple(headings),
                    locator_type="paragraph",
                    locator_start=paragraph_index,
                    locator_end=paragraph_index,
                )
            )
    return collector.result()


def _parse_pptx(path: Path, limits: ParserLimits) -> ParsedDocument:
    collector = _BlockCollector(limits)
    presentation = Presentation(path)
    for slide_index, slide in enumerate(presentation.slides, 1):
        parts: list[str] = []
        title = slide.shapes.title
        title_text = title.text.strip() if title is not None else ""
        if title_text:
            parts.append(title_text)
        for shape in slide.shapes:
            if shape is title:
                continue
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    value = " | ".join(cell.text.strip() for cell in row.cells)
                    if value.strip(" |"):
                        parts.append(value)
            elif getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    parts.append(value)
        collector.add(
            StructuredBlock(
                text="\n".join(parts),
                heading_path=(title_text,) if title_text else (),
                locator_type="slide",
                locator_start=slide_index,
                locator_end=slide_index,
            )
        )
    return collector.result(slide_count=len(presentation.slides))


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_xlsx(path: Path, limits: ParserLimits) -> ParsedDocument:
    collector = _BlockCollector(limits)
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if len(workbook.worksheets) > limits.max_sheets:
            raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
        nonempty_cells = 0
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [_cell_text(value) for value in row]
                nonempty_cells += sum(bool(value) for value in values)
                if nonempty_cells > limits.max_cells:
                    raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
                if not any(values):
                    continue
                collector.add(
                    StructuredBlock(
                        text=" | ".join(values),
                        heading_path=(sheet.title,),
                        locator_type="sheet_rows",
                        locator_start=row_index,
                        locator_end=row_index,
                        sheet_name=sheet.title,
                    )
                )
        return collector.result(sheet_count=len(workbook.worksheets))
    finally:
        workbook.close()


def _decode_text(path: Path, limits: ParserLimits) -> str:
    raw = path.read_bytes()
    if b"\0" in raw:
        raise KnowledgeParseError("KNOWLEDGE_PARSE_FAILED")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            value = raw.decode(encoding, errors="strict")
            break
        except UnicodeDecodeError:
            continue
    else:
        raise KnowledgeParseError("KNOWLEDGE_PARSE_FAILED")
    if len(value) > limits.max_text_characters:
        raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
    return value


def _parse_csv(path: Path, limits: ParserLimits) -> ParsedDocument:
    collector = _BlockCollector(limits)
    content = _decode_text(path, limits)
    reader = csv.reader(content.splitlines())
    try:
        for row_index, row in enumerate(reader, 1):
            if row_index > limits.max_csv_rows:
                raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
            if len(row) > limits.max_csv_columns:
                raise KnowledgeParseError("KNOWLEDGE_PARSE_LIMIT_EXCEEDED")
            values = [value.strip() for value in row]
            if not any(values):
                continue
            collector.add(
                StructuredBlock(
                    text=" | ".join(values),
                    heading_path=(path.stem,),
                    locator_type="sheet_rows",
                    locator_start=row_index,
                    locator_end=row_index,
                    sheet_name=path.stem[:128],
                )
            )
    except csv.Error as exc:
        raise KnowledgeParseError("KNOWLEDGE_PARSE_FAILED") from exc
    return collector.result(sheet_count=1)


def _text_paragraphs(value: str) -> list[str]:
    return [
        paragraph.strip()
        for paragraph in re.split(r"(?:\r?\n){2,}", value)
        if paragraph.strip()
    ]


def _parse_text(path: Path, limits: ParserLimits) -> ParsedDocument:
    collector = _BlockCollector(limits)
    content = _decode_text(path, limits)
    for index, paragraph in enumerate(_text_paragraphs(content), 1):
        collector.add(
            StructuredBlock(
                text=paragraph,
                heading_path=(),
                locator_type="paragraph",
                locator_start=index,
                locator_end=index,
            )
        )
    return collector.result()


def _parse_markdown(path: Path, limits: ParserLimits) -> ParsedDocument:
    collector = _BlockCollector(limits)
    content = _decode_text(path, limits)
    headings: list[str] = []
    paragraph_lines: list[str] = []
    paragraph_index = 0

    def flush() -> None:
        nonlocal paragraph_index
        text = "\n".join(paragraph_lines).strip()
        paragraph_lines.clear()
        if not text:
            return
        paragraph_index += 1
        collector.add(
            StructuredBlock(
                text=text,
                heading_path=tuple(headings),
                locator_type="paragraph",
                locator_start=paragraph_index,
                locator_end=paragraph_index,
            )
        )

    for line in content.splitlines():
        matched = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if matched:
            flush()
            level = len(matched.group(1))
            headings = headings[: level - 1]
            headings.append(matched.group(2).strip())
        elif line.strip():
            paragraph_lines.append(line)
        else:
            flush()
    flush()
    return collector.result()


def parse_document(
    path: Path,
    limits: ParserLimits,
    *,
    extension: str | None = None,
) -> ParsedDocument:
    file_path = Path(path)
    selected_extension = (extension or file_path.suffix).lower()
    if selected_extension not in _SUPPORTED_EXTENSIONS:
        raise KnowledgeParseError("KNOWLEDGE_FORMAT_UNSUPPORTED")
    if not file_path.is_file():
        raise KnowledgeParseError("KNOWLEDGE_OBJECT_MISSING")
    _validate_signature(file_path, selected_extension)
    if selected_extension in _OFFICE_EXTENSIONS:
        _validate_office_archive(file_path, limits)
    parser = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".pptx": _parse_pptx,
        ".xlsx": _parse_xlsx,
        ".csv": _parse_csv,
        ".txt": _parse_text,
        ".md": _parse_markdown,
        ".markdown": _parse_markdown,
    }[selected_extension]
    try:
        return parser(file_path, limits)
    except KnowledgeParseError:
        raise
    except Exception as exc:
        raise KnowledgeParseError("KNOWLEDGE_PARSE_FAILED") from exc
